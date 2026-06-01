"""Informe de integridad NoofitPro ↔ cache local.

Calcula EN VIVO (sin tabla local de reservas) el diff entre los clientes que
están reservando en clases del manager actual (en NoofitPro) y los que figuran
en `cliente_cache`. Sirve para detectar "fantasmas" — clientes que reservan
pero no están dados de alta en el centro / no aparecen en la web admin.

Endpoints:
  GET /api/integridad/reservas-sin-cliente?dias=90       JSON
  GET /api/integridad/reservas-sin-cliente/excel?dias=90 XLSX descarga

NO persistimos las reservas localmente — se consultan a NoofitPro en cada
llamada. Si el reporte se vuelve lento podemos cachear N minutos en memoria
del proceso, pero NUNCA en tabla.
"""
import io
import logging
from datetime import date, timedelta
from collections import Counter
from flask import Blueprint, request, jsonify, g, send_file

from ..auth import auth_required
from ..db import get_conn
from .. import noofit_client as nc

bp = Blueprint('integridad', __name__)
log = logging.getLogger(__name__)


def _trainers_de_manager(id_manager):
    """Devuelve el conjunto de id_trainer asociados al manager (los que
    aparecen en cliente_cache.id_trainer + el propio manager por si coincide
    con un trainer)."""
    with get_conn() as conn, conn.cursor() as cur:
        cur.execute(
            "SELECT DISTINCT id_trainer FROM cliente_cache "
            "WHERE id_manager=%s AND id_trainer IS NOT NULL",
            (str(id_manager),))
        ts = {r['id_trainer'] for r in cur.fetchall() if r['id_trainer']}
    # Incluir el id_manager por si actúa también como trainer (caso Round)
    ts.add(str(id_manager))
    # Convertir a ints donde se pueda (NoofitPro usa ints)
    out = set()
    for t in ts:
        try: out.add(int(t))
        except (TypeError, ValueError): pass
    return out


def _calcular_fantasmas(id_manager, dias):
    """Calcula el diff en vivo. Devuelve dict con totales + lista de fantasmas."""
    id_manager = str(id_manager)

    # 1) Cache local — qué clientes vemos
    with get_conn() as conn, conn.cursor() as cur:
        cur.execute(
            "SELECT id, raw_data->>'name' AS nombre, raw_data->>'surname' AS apellidos, "
            "raw_data->>'email' AS email, raw_data->>'dni' AS dni "
            "FROM cliente_cache WHERE id_manager=%s",
            (id_manager,))
        cache_rows = cur.fetchall()
    cache_ids = {int(r['id']) for r in cache_rows}

    # 2) Trainers de este manager (para filtrar las salas)
    trainers = _trainers_de_manager(id_manager)

    # 3) Reservas confirmadas EN VIVO (NoofitPro)
    hoy = date.today()
    desde = (hoy - timedelta(days=int(dias))).isoformat() + 'T00:00:00+02:00'
    hasta = (hoy + timedelta(days=1)).isoformat() + 'T00:00:00+02:00'
    reservas = nc.get_reservas_confirmadas(desde, hasta) or []
    # Filtrar salas a los trainers del manager
    reservas_mgr = [r for r in reservas if r.get('id_trainer') in trainers]

    # 4) Agrupar por cliente
    por_cliente = {}
    for r in reservas_mgr:
        try: cid = int(r['cliente_id'])
        except (TypeError, ValueError): continue
        por_cliente.setdefault(cid, []).append(r)

    # 5) Fantasmas = reservan pero no en cache
    fantasmas_ids = sorted(set(por_cliente) - cache_ids)

    # 6) Cross-manager: ¿en cache de OTRO manager?
    cross_map = {}
    if fantasmas_ids:
        placeholders = ','.join(['%s'] * len(fantasmas_ids))
        with get_conn() as conn, conn.cursor() as cur:
            cur.execute(
                f"SELECT id, id_manager, raw_data->>'name' AS n, "
                f"raw_data->>'surname' AS s, raw_data->>'email' AS e, "
                f"raw_data->>'dni' AS d "
                f"FROM cliente_cache WHERE id IN ({placeholders})",
                list(fantasmas_ids))
            for r in cur.fetchall():
                cross_map.setdefault(int(r['id']), []).append({
                    'id_manager': r['id_manager'],
                    'nombre': (r['n'] or '') + ' ' + (r['s'] or ''),
                    'email': r['e'], 'dni': r['d'],
                })

    # 7) Construir filas finales
    fantasmas = []
    for cid in fantasmas_ids:
        rs = por_cliente[cid]
        fechas = sorted({r.get('fecha') for r in rs if r.get('fecha')})
        actividades = sorted({r.get('actividad_nombre') for r in rs if r.get('actividad_nombre')})
        cruce = cross_map.get(cid, [])
        # Nombre fallback desde la reserva (nameClient viene de NoofitPro
        # incluso si el cliente pertenece a otro manager/cuenta).
        nombre_reserva = next((r.get('cliente_nombre') for r in rs if r.get('cliente_nombre')), None)
        nombre_cache = cruce[0]['nombre'].strip() if cruce else None
        fantasmas.append({
            'cliente_id': cid,
            'reservas': len(rs),
            'primera_fecha': fechas[0] if fechas else None,
            'ultima_fecha':  fechas[-1] if fechas else None,
            'actividades': actividades,
            'en_otros_managers': cruce,  # [] si es verdadero fantasma
            'tipo': ('cross_manager' if cruce else 'verdadero_fantasma'),
            'nombre': nombre_cache or nombre_reserva,
            'email':  cruce[0]['email'] if cruce else None,
            'dni':    cruce[0]['dni'] if cruce else None,
        })

    # Inactivos en cache (en cache pero sin reservas en el rango)
    inactivos = sorted(cache_ids - set(por_cliente))

    return {
        'id_manager': id_manager,
        'dias': int(dias),
        'cache_total': len(cache_ids),
        'reservas_total': len(reservas_mgr),
        'clientes_reservando': len(por_cliente),
        'fantasmas_total': len(fantasmas),
        'fantasmas_verdaderos': sum(1 for f in fantasmas if f['tipo'] == 'verdadero_fantasma'),
        'fantasmas_cross_manager': sum(1 for f in fantasmas if f['tipo'] == 'cross_manager'),
        'inactivos_en_cache': len(inactivos),
        'fantasmas': fantasmas,
    }


@bp.route('/reservas-sin-cliente', methods=['GET'])
@auth_required
def reservas_sin_cliente():
    """Diff EN VIVO (NoofitPro vs cache local) — JSON."""
    dias = request.args.get('dias', '90')
    try: dias = max(1, min(365, int(dias)))
    except ValueError: dias = 90
    try:
        data = _calcular_fantasmas(g.id_manager, dias)
        return jsonify({'ok': True, **data})
    except Exception as e:
        log.exception('reservas_sin_cliente')
        return jsonify({'ok': False, 'error': str(e)}), 500


@bp.route('/reservas-sin-cliente/excel', methods=['GET'])
@auth_required
def reservas_sin_cliente_excel():
    """Mismo diff en formato Excel descargable."""
    try:
        from openpyxl import Workbook
    except ImportError:
        return jsonify({'ok': False, 'error': 'openpyxl_no_disponible'}), 500

    dias = request.args.get('dias', '90')
    try: dias = max(1, min(365, int(dias)))
    except ValueError: dias = 90

    data = _calcular_fantasmas(g.id_manager, dias)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Fantasmas'
    headers = ['Tipo', 'cliente_id', 'Nombre', 'Email', 'DNI',
               '#Reservas', 'Primera', 'Última', 'Actividades',
               'En otros managers']
    ws.append(headers)
    for f in data['fantasmas']:
        ws.append([
            f['tipo'], f['cliente_id'],
            f.get('nombre') or '', f.get('email') or '', f.get('dni') or '',
            f['reservas'], f.get('primera_fecha') or '', f.get('ultima_fecha') or '',
            ', '.join(f.get('actividades') or []),
            ', '.join(o.get('id_manager') or '' for o in f.get('en_otros_managers') or []),
        ])
    # Ajustar anchos básicos
    widths = [22, 14, 28, 30, 14, 10, 12, 12, 50, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    # Resumen en una segunda hoja
    ws2 = wb.create_sheet('Resumen')
    ws2.append(['Manager', data['id_manager']])
    ws2.append(['Días analizados', data['dias']])
    ws2.append(['Clientes en cache local', data['cache_total']])
    ws2.append(['Reservas en el rango', data['reservas_total']])
    ws2.append(['Clientes únicos reservando', data['clientes_reservando']])
    ws2.append(['Fantasmas (total)', data['fantasmas_total']])
    ws2.append(['  Verdaderos (sin cache en ningún manager)', data['fantasmas_verdaderos']])
    ws2.append(['  Cross-manager (en cache de otro manager)', data['fantasmas_cross_manager']])
    ws2.append(['Clientes en cache sin reservas', data['inactivos_en_cache']])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f'integridad_reservas_{g.id_manager}_{date.today().isoformat()}.xlsx'
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
